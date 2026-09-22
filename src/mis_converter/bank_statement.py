"""Bank Statement PDF Parser - Meezan Bank Format.

Extracts transactions from Meezan Bank statement PDFs and generates a formatted
Excel report with opening balance, closing balance, transaction details,
and totals for credits and debits.

SUPPORTED FORMAT: Meezan Bank online account statements
FORMAT VERSION: As of 2025-2026 statement layout

To add support for other banks, extend the _parse_header() and
_parse_transactions() functions with bank-specific logic.
"""

from __future__ import annotations

import re
from datetime import datetime
from typing import BinaryIO

import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


def extract_transactions(pdf_path: str | BinaryIO) -> dict:
    """Extract transactions and metadata from bank statement PDF.

    Returns
    -------
    dict with keys:
        - account_title: str
        - account_number: str
        - iban: str
        - currency: str
        - from_date: str
        - to_date: str
        - opening_balance: float
        - closing_balance: float
        - transactions: list of dicts with keys:
            - date: str
            - description: str
            - credit: float
            - debit: float
            - balance: float
    """
    with pdfplumber.open(pdf_path) as pdf:
        # Extract metadata from first page
        first_page = pdf.pages[0]
        first_text = first_page.extract_text()

        metadata = _parse_header(first_text)

        # Extract transactions from all pages
        transactions = []
        for page in pdf.pages:
            page_text = page.extract_text()
            page_transactions = _parse_transactions(page_text)
            transactions.extend(page_transactions)

        metadata['transactions'] = transactions
        return metadata


def _parse_header(text: str) -> dict:
    """Parse header information from first page."""
    lines = text.split('\n')

    result = {
        'account_title': '',
        'account_number': '',
        'iban': '',
        'currency': '',
        'from_date': '',
        'to_date': '',
        'opening_balance': 0.0,
        'closing_balance': 0.0,
    }

    # Find account info (line with account title and numbers)
    for i, line in enumerate(lines):
        if 'Account Title' in line and i + 1 < len(lines):
            # Next line has: ACCOUNT_NAME ACCOUNT_NUMBER IBAN
            parts = lines[i + 1].split()
            if len(parts) >= 3:
                # Account title is everything except last 2 items (number and IBAN)
                result['account_title'] = ' '.join(parts[:-2])
                result['account_number'] = parts[-2]
                result['iban'] = parts[-1]

    # Find currency and date range
    for i, line in enumerate(lines):
        if 'Currency' in line and 'From Date' in line and i + 1 < len(lines):
            # Next line has: Pakistan Rupee(PKR) 01 Jul 2025 30 Jun 2026
            next_line = lines[i + 1]
            match = re.search(r'\(([A-Z]+)\)', next_line)
            if match:
                result['currency'] = match.group(1)

            # Extract dates
            date_match = re.findall(r'\d{2}\s+[A-Za-z]+\s+\d{4}', next_line)
            if len(date_match) >= 2:
                result['from_date'] = date_match[0]
                result['to_date'] = date_match[1]

    # Find opening and closing balance
    for i, line in enumerate(lines):
        if 'Opening Balance' in line and 'Closing Balance' in line:
            # Next line has balances
            if i + 1 < len(lines):
                balance_line = lines[i + 1]
                # Extract all PKR amounts
                amounts = re.findall(r'PKR([\d,]+\.\d{2})', balance_line)
                if len(amounts) >= 2:
                    result['opening_balance'] = float(amounts[0].replace(',', ''))
                    result['closing_balance'] = float(amounts[1].replace(',', ''))

    return result


def _parse_transactions(text: str) -> list[dict]:
    """Parse transactions from page text.

    Transaction format:
    02 Jul 2025 KABABJEES EXPRESS - POS - PKR2,425.00 PKR9,908,113.15
                Transaction STAN (109350)

    Or with credit:
    01 Jul 2025 Payment of Profit + PKR51,240.38 PKR9,957,254.53
    """
    transactions = []
    lines = text.split('\n')

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Match transaction line starting with date
        date_match = re.match(r'^(\d{2}\s+[A-Za-z]+\s+\d{4})\s+(.+)$', line)

        # Skip header lines. A line starting with a booking date is always a
        # transaction, even when its description contains the words
        # "Credit"/"Debit" (e.g. "Batch Transfer - Credit SALARY TRF").
        if date_match is None and any(
            x in line for x in ['Account Statement', 'Booking Date', 'Description', 'Credit', 'Debit']
        ):
            i += 1
            continue

        if date_match:
            date_str = date_match.group(1)
            rest = date_match.group(2).strip()

            # Check if next line is continuation (doesn't start with date)
            description_parts = []
            if rest:
                description_parts.append(rest)

            # Look ahead for continuation lines
            j = i + 1
            while j < len(lines):
                next_line = lines[j].strip()
                if not next_line:
                    j += 1
                    break
                # If next line starts with a date, it's a new transaction
                if re.match(r'^\d{2}\s+[A-Za-z]+\s+\d{4}', next_line):
                    break
                # If it's a number-only line or footer, stop
                if re.match(r'^\d+\s+\d{2}\s+[A-Za-z]+\s+\d{4}', next_line):
                    break
                description_parts.append(next_line)
                j += 1

            # Combine description
            full_text = ' '.join(description_parts)

            # Extract amounts and determine credit/debit
            # Pattern: [+-] PKR amount at the end
            credit = 0.0
            debit = 0.0
            balance = 0.0
            description = full_text

            # Find all PKR amounts (must include PKR prefix)
            amount_pattern = r'([+-])?\s*PKR([\d,]+\.\d{2})'
            amounts = re.findall(amount_pattern, full_text)

            if amounts:
                # Last amount is usually the balance
                if len(amounts) >= 2:
                    # Second to last is the transaction amount
                    sign, amount_str = amounts[-2]
                    amount = float(amount_str.replace(',', ''))
                    if sign == '+':
                        credit = amount
                    elif sign == '-':
                        debit = amount

                    # Last is the balance
                    balance = float(amounts[-1][1].replace(',', ''))

                    # Remove amounts from description
                    description = re.sub(r'[+-]?\s*PKR[\d,]+\.\d{2}', '', full_text).strip()
                elif len(amounts) == 1:
                    # Only balance (shouldn't happen normally)
                    balance = float(amounts[0][1].replace(',', ''))
                    description = re.sub(r'[+-]?\s*PKR[\d,]+\.\d{2}', '', full_text).strip()

            if description:  # Only add if we have a description
                transactions.append({
                    'date': date_str,
                    'description': description,
                    'credit': credit,
                    'debit': debit,
                    'balance': balance,
                })

            i = j
        else:
            i += 1

    return transactions


def write_excel(data: dict, output_path: str | BinaryIO) -> None:
    """Write bank statement data to formatted Excel file.

    Layout:
        Row 1: "BANK STATEMENT" title
        Row 2: Account details
        Row 3: Date range
        Row 4: Opening balance
        Row 5: Blank
        Row 6: Column headers
        Row 7+: Transactions
        Last rows: Totals and closing balance
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Statement"

    # Styles
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11)
    data_font = Font(size=10)
    total_font = Font(bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    accounting_format = '#,##0.00'

    thick_side = Side(style="medium")
    thin_side = Side(style="thin")
    outer_border = Border(
        top=thick_side, bottom=thick_side,
        left=thick_side, right=thick_side,
    )

    current_row = 1

    # Row 1: Title
    ws.merge_cells(f'A{current_row}:E{current_row}')
    cell = ws.cell(row=current_row, column=1, value="BANK STATEMENT")
    cell.font = title_font
    cell.alignment = center_align
    for col_idx in range(1, 6):
        ws.cell(row=current_row, column=col_idx).border = outer_border
    current_row += 1

    # Row 2: Account details
    account_info = f"{data['account_title']} | {data['account_number']} | {data['iban']}"
    ws.merge_cells(f'A{current_row}:E{current_row}')
    cell = ws.cell(row=current_row, column=1, value=account_info)
    cell.font = Font(size=10)
    cell.alignment = center_align
    current_row += 1

    # Row 3: Date range
    date_range = f"Period: {data['from_date']} to {data['to_date']}"
    ws.merge_cells(f'A{current_row}:E{current_row}')
    cell = ws.cell(row=current_row, column=1, value=date_range)
    cell.font = Font(size=10)
    cell.alignment = center_align
    current_row += 1

    # Row 4: Opening balance
    ws.cell(row=current_row, column=1, value="Opening Balance:").font = Font(bold=True, size=10)
    cell = ws.cell(row=current_row, column=2, value=data['opening_balance'])
    cell.number_format = accounting_format
    cell.font = Font(bold=True, size=10)
    current_row += 1

    # Blank row
    current_row += 1

    # Column headers
    headers = ['Date', 'Description', 'Credit', 'Debit', 'Balance']
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center_align
        left = thick_side if col_idx == 1 else thin_side
        right = thick_side if col_idx == len(headers) else thin_side
        cell.border = Border(
            top=thick_side, bottom=thick_side,
            left=left, right=right,
        )
    current_row += 1

    # Transaction rows
    transactions = data['transactions']
    for txn in transactions:
        ws.cell(row=current_row, column=1, value=txn['date']).font = data_font
        ws.cell(row=current_row, column=2, value=txn['description']).font = data_font

        credit_cell = ws.cell(row=current_row, column=3, value=txn['credit'] if txn['credit'] > 0 else '')
        credit_cell.font = data_font
        credit_cell.alignment = center_align
        if txn['credit'] > 0:
            credit_cell.number_format = accounting_format

        debit_cell = ws.cell(row=current_row, column=4, value=txn['debit'] if txn['debit'] > 0 else '')
        debit_cell.font = data_font
        debit_cell.alignment = center_align
        if txn['debit'] > 0:
            debit_cell.number_format = accounting_format

        balance_cell = ws.cell(row=current_row, column=5, value=txn['balance'])
        balance_cell.font = data_font
        balance_cell.alignment = center_align
        balance_cell.number_format = accounting_format

        current_row += 1

    # Blank row
    current_row += 1

    # Totals
    total_credits = sum(t['credit'] for t in transactions)
    total_debits = sum(t['debit'] for t in transactions)

    # Total Credits
    ws.cell(row=current_row, column=2, value="Total Credits:").font = total_font
    cell = ws.cell(row=current_row, column=3, value=total_credits)
    cell.font = total_font
    cell.number_format = accounting_format
    cell.alignment = center_align
    current_row += 1

    # Total Debits
    ws.cell(row=current_row, column=2, value="Total Debits:").font = total_font
    cell = ws.cell(row=current_row, column=4, value=total_debits)
    cell.font = total_font
    cell.number_format = accounting_format
    cell.alignment = center_align
    current_row += 1

    # Closing Balance
    ws.cell(row=current_row, column=2, value="Closing Balance:").font = total_font
    cell = ws.cell(row=current_row, column=5, value=data['closing_balance'])
    cell.font = total_font
    cell.number_format = accounting_format
    cell.alignment = center_align
    current_row += 1

    # Column widths
    ws.column_dimensions['A'].width = 15  # Date
    ws.column_dimensions['B'].width = 50  # Description
    ws.column_dimensions['C'].width = 18  # Credit
    ws.column_dimensions['D'].width = 18  # Debit
    ws.column_dimensions['E'].width = 20  # Balance

    wb.save(output_path)


def convert(input_path: str | BinaryIO, output_path: str | BinaryIO) -> dict:
    """Main orchestrator: extract → write Excel.

    Returns the extracted data dict for display.
    """
    data = extract_transactions(input_path)
    write_excel(data, output_path)
    return data
